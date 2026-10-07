#!/usr/bin/env python3
"""
anonymise.py - folder-based PII anonymisation for .csv and .docx files.

Pipeline (per text unit):
  1. Presidio Analyzer   - built-in recognisers + spaCy NER
                         - extra regex recognisers from settings.json (e.g. NHS number)
                         - keyword recognisers for ORGANIZATION and ORGANIZATION_RELATED_ENTITY
  2. fastcoref           - coreference clusters; any mention that co-refers with a detected
                           entity ("Mr Smith", "Smith", "the Acme group") gets the same placeholder
  3. Presidio Anonymizer - replaces each span with a session-consistent placeholder
                           (same PII across all files in a session -> same value)
  4. mapping.json        - what was replaced with what, per placeholder and per file

Usage:
  python anonymise.py --input ./in --output ./out --settings settings.json
  python anonymise.py --input ./in --output ./out --resume-mapping ./out/mapping.json   # continue a session
"""
from __future__ import annotations

import argparse
import bisect
import csv
import json
import logging
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Optional

from presidio_analyzer import (
    AnalyzerEngine,
    EntityRecognizer,
    Pattern,
    PatternRecognizer,
    RecognizerResult,
)
from presidio_analyzer.nlp_engine import NlpEngineProvider
from presidio_anonymizer import AnonymizerEngine
from presidio_anonymizer.entities import OperatorConfig
from presidio_anonymizer.entities import RecognizerResult as AnonymizerResult

log = logging.getLogger("anonymise")

SUPPORTED_SUFFIXES = {".csv", ".docx"}


# --------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------
def norm(text: str) -> str:
    """Normalise a surface string so trivially different spellings share one placeholder."""
    return " ".join(text.split()).casefold()


@dataclass
class Span:
    start: int
    end: int
    entity_type: str
    score: float
    source: str  # keyword | regex | presidio | coref
    text: str


SOURCE_PRIORITY = {"keyword": 3, "regex": 2, "presidio": 1, "coref": 0}


def resolve_overlaps(spans: list[Span]) -> list[Span]:
    """Greedy non-overlapping selection: keyword lists > custom regex > Presidio > coref,
    then by score, then by length."""
    ordered = sorted(
        spans,
        key=lambda s: (-SOURCE_PRIORITY.get(s.source, 0), -s.score, -(s.end - s.start), s.start),
    )
    chosen: list[Span] = []
    for s in ordered:
        if all(s.end <= c.start or s.start >= c.end for c in chosen):
            chosen.append(s)
    return sorted(chosen, key=lambda s: s.start)


# --------------------------------------------------------------------------------------
# Validators usable from settings.json ("validator": "nhs_mod11")
# --------------------------------------------------------------------------------------
def nhs_mod11(value: str) -> bool:
    digits = re.sub(r"\D", "", value)
    if len(digits) != 10:
        return False
    total = sum(int(d) * w for d, w in zip(digits[:9], range(10, 1, -1)))
    check = 11 - (total % 11)
    if check == 11:
        check = 0
    if check == 10:
        return False
    return check == int(digits[9])


def luhn(value: str) -> bool:
    digits = [int(d) for d in re.sub(r"\D", "", value)]
    if len(digits) < 12:
        return False
    checksum = 0
    for i, d in enumerate(reversed(digits)):
        if i % 2:
            d *= 2
            if d > 9:
                d -= 9
        checksum += d
    return checksum % 10 == 0


VALIDATORS: dict[str, Callable[[str], bool]] = {"nhs_mod11": nhs_mod11, "luhn": luhn}


# --------------------------------------------------------------------------------------
# Custom recognisers
# --------------------------------------------------------------------------------------
class ConfigurableRegexRecognizer(PatternRecognizer):
    """PatternRecognizer built from settings.json, with an optional checksum validator."""

    def __init__(self, cfg: dict, language: str):
        patterns = [
            Pattern(name=p.get("name", cfg["entity"]), regex=p["regex"], score=float(p.get("score", 0.6)))
            for p in cfg["patterns"]
        ]
        super().__init__(
            supported_entity=cfg["entity"],
            name=f"Regex:{cfg.get('name', cfg['entity'])}",
            patterns=patterns,
            context=cfg.get("context") or None,
            supported_language=language,
        )
        v = cfg.get("validator")
        if v and v not in VALIDATORS:
            raise ValueError(f"Unknown validator '{v}' for {cfg['entity']}. Options: {list(VALIDATORS)}")
        self._validator = VALIDATORS.get(v) if v else None

    def validate_result(self, pattern_text: str) -> Optional[bool]:
        # True -> score 1.0, False -> discarded, None -> keep pattern score
        return self._validator(pattern_text) if self._validator else None


class KeywordRecognizer(EntityRecognizer):
    """Dictionary recogniser for long / messy names (organisations, products).

    Entries are either "Name" or {"name": "Canonical Ltd", "aliases": ["Canon", "CL Group"]}.
    All aliases resolve to the canonical name, so they share one placeholder.
    Matching is case-insensitive, whitespace/line-break tolerant, longest-match-first,
    and respects word boundaries.
    """

    def __init__(self, entity: str, entries: list, language: str, score: float = 1.0,
                 case_sensitive: bool = False, name: Optional[str] = None):
        self.entity = entity
        self.score = score
        self.alias_to_canonical: dict[str, str] = {}
        for entry in entries:
            if isinstance(entry, str):
                canonical, aliases = entry, []
            else:
                canonical, aliases = entry["name"], entry.get("aliases", [])
            for alias in [canonical, *aliases]:
                if alias and alias.strip():
                    self.alias_to_canonical[norm(alias)] = canonical

        alternatives = []
        for alias in sorted(self.alias_to_canonical, key=len, reverse=True):
            tokens = alias.split(" ")
            alternatives.append(r"\s+".join(re.escape(t) for t in tokens))
        flags = 0 if case_sensitive else re.IGNORECASE
        self.pattern = (
            # boundaries also exclude '@' / '.x' so we never match inside emails, URLs or domains
            re.compile(r"(?<![\w@.])(?:" + "|".join(alternatives) + r")(?![\w@]|\.\w)", flags)
            if alternatives else None
        )
        super().__init__(supported_entities=[entity], name=name or f"Keyword:{entity}",
                         supported_language=language)

    def load(self) -> None:  # nothing to load
        pass

    def canonical(self, text: str) -> str:
        return self.alias_to_canonical.get(norm(text), text)

    def analyze(self, text, entities, nlp_artifacts=None):
        if not self.pattern or (entities and self.entity not in entities):
            return []
        results = []
        for m in self.pattern.finditer(text):
            results.append(
                RecognizerResult(
                    entity_type=self.entity,
                    start=m.start(),
                    end=m.end(),
                    score=self.score,
                    recognition_metadata={
                        RecognizerResult.RECOGNIZER_NAME_KEY: self.name,
                        RecognizerResult.RECOGNIZER_IDENTIFIER_KEY: self.id,
                        "canonical": self.canonical(m.group(0)),
                    },
                )
            )
        return results


def load_keyword_entries(cfg: dict, base_dir: Path) -> list:
    """Inline 'keywords' plus an optional 'keywords_file' (one entry per line,
    aliases separated by '|', e.g.  Acme Holdings Ltd|Acme|AHL )."""
    entries = list(cfg.get("keywords", []))
    kw_file = cfg.get("keywords_file")
    if kw_file:
        path = (base_dir / kw_file) if not Path(kw_file).is_absolute() else Path(kw_file)
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = [p.strip() for p in line.split("|") if p.strip()]
            entries.append({"name": parts[0], "aliases": parts[1:]})
    return entries


# --------------------------------------------------------------------------------------
# Coreference (fastcoref)
# --------------------------------------------------------------------------------------
PRONOUN_LIKE = {
    "i", "me", "my", "mine", "myself", "we", "us", "our", "ours", "ourselves", "you", "your",
    "yours", "yourself", "he", "him", "his", "himself", "she", "her", "hers", "herself", "it",
    "its", "itself", "they", "them", "their", "theirs", "themselves", "this", "that", "these",
    "those", "the", "a", "an", "mr", "mrs", "ms", "miss", "dr", "prof", "sir", "dame", "who",
    "whom", "whose", "which",
}


def looks_like_name(mention: str) -> bool:
    """True if the mention contains a capitalised token that isn't a pronoun/determiner/title.
    'Mr Smith' -> True, 'he' -> False, 'the company' -> False."""
    for tok in re.findall(r"[^\W\d_][\w'\-]*", mention):
        if tok[0].isupper() and tok.casefold().rstrip("'s") not in PRONOUN_LIKE:
            return True
    return False


class CorefResolver:
    def __init__(self, cfg: dict):
        from fastcoref import FCoref, LingMessCoref  # imported lazily: heavy

        model_cls = LingMessCoref if str(cfg.get("model", "FCoref")).lower() == "lingmesscoref" else FCoref
        kwargs = {"device": cfg.get("device", "cpu")}
        if cfg.get("model_name_or_path"):
            kwargs["model_name_or_path"] = cfg["model_name_or_path"]
        if cfg.get("spacy_model"):
            kwargs["nlp"] = cfg["spacy_model"]
        try:
            self.model = model_cls(enable_progress_bar=False, **kwargs)
        except TypeError:  # older signatures without enable_progress_bar
            self.model = model_cls(**kwargs)
        self.max_chars = int(cfg.get("max_chars_per_chunk", 5000))

    def _chunks(self, text: str) -> list[tuple[int, str]]:
        """Split long text on line breaks into chunks of <= max_chars, keeping offsets."""
        if len(text) <= self.max_chars:
            return [(0, text)]
        chunks, start, cur = [], 0, 0
        for m in re.finditer(r"\n", text):
            if m.end() - start > self.max_chars and cur > start:
                chunks.append((start, text[start:cur]))
                start = cur
            cur = m.end()
        chunks.append((start, text[start:]))
        return chunks

    def clusters(self, text: str) -> list[list[tuple[int, int]]]:
        chunks = [(off, t) for off, t in self._chunks(text) if t.strip()]
        if not chunks:
            return []
        preds = self.model.predict(texts=[t for _, t in chunks])
        out = []
        for (off, _), pred in zip(chunks, preds):
            for cluster in pred.get_clusters(as_strings=False):
                out.append([(s + off, e + off) for s, e in cluster])
        return out


# --------------------------------------------------------------------------------------
# Session: consistent replacements across all files
# --------------------------------------------------------------------------------------
class Session:
    def __init__(self, replacement_format: str, alias_scope: str = "session"):
        self.fmt = replacement_format
        self.alias_scope = alias_scope  # where coref-derived aliases live: "session" | "file"
        self.counters: dict[str, int] = defaultdict(int)
        self.placeholders: dict[tuple[str, str], str] = {}        # (entity, canonical) -> placeholder
        self.aliases: dict[tuple[str, str], str] = {}             # (entity, alias) -> canonical
        self.file_aliases: dict[tuple[str, str], str] = {}
        self.records: dict[str, dict] = {}                        # placeholder -> info
        self.file_log: dict[str, list[dict]] = defaultdict(list)
        self.files_processed: list[str] = []
        self.current_file: Optional[str] = None

    # -- aliases ------------------------------------------------------------------------
    def canonical(self, entity: str, text: str) -> str:
        key = (entity, norm(text))
        return self.file_aliases.get(key) or self.aliases.get(key) or key[1]

    def add_alias(self, entity: str, alias: str, canonical_text: str, scope: str = "session") -> None:
        target = self.canonical(entity, canonical_text)
        key = (entity, norm(alias))
        if key[1] == target:
            return
        (self.aliases if scope == "session" else self.file_aliases)[key] = target

    # -- replacements -------------------------------------------------------------------
    def replacement(self, entity: str, original: str) -> str:
        canon = self.canonical(entity, original)
        key = (entity, canon)
        if key not in self.placeholders:
            self.counters[entity] += 1
            ph = self.fmt.format(entity_type=entity, index=self.counters[entity])
            self.placeholders[key] = ph
            self.records[ph] = {"entity_type": entity, "canonical": canon, "originals": []}
        ph = self.placeholders[key]
        if original not in self.records[ph]["originals"]:
            self.records[ph]["originals"].append(original)
        return ph

    def start_file(self, name: str) -> None:
        self.current_file = name
        self.file_aliases = {}
        if name not in self.files_processed:
            self.files_processed.append(name)

    def log(self, span: Span, replacement: str, location: str = "") -> None:
        entry = {
            "entity_type": span.entity_type,
            "original": span.text,
            "replacement": replacement,
            "source": span.source,
            "score": round(span.score, 3),
            "start": span.start,
            "end": span.end,
        }
        if location:
            entry["location"] = location
        self.file_log[self.current_file].append(entry)

    # -- persistence --------------------------------------------------------------------
    def to_json(self, input_dir: str, settings_path: str) -> dict:
        lookup: dict[str, dict[str, str]] = defaultdict(dict)
        for (entity, canon), ph in self.placeholders.items():
            lookup[entity][canon] = ph
        for (entity, alias), canon in self.aliases.items():
            if (entity, canon) in self.placeholders:
                lookup[entity][alias] = self.placeholders[(entity, canon)]
        return {
            "session": {
                "updated": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "input_dir": input_dir,
                "settings": settings_path,
                "replacement_format": self.fmt,
                "files_processed": self.files_processed,
            },
            "counters": dict(self.counters),
            "replacements": self.records,           # placeholder -> what it stands for
            "lookup": {k: dict(v) for k, v in lookup.items()},  # entity -> normalised text -> placeholder
            "aliases": [
                {"entity_type": e, "alias": a, "canonical": c} for (e, a), c in self.aliases.items()
            ],
            "files": dict(self.file_log),           # every individual replacement
        }

    @classmethod
    def from_json(cls, data: dict, replacement_format: str, alias_scope: str) -> "Session":
        s = cls(data.get("session", {}).get("replacement_format", replacement_format), alias_scope)
        s.counters.update(data.get("counters", {}))
        for ph, info in data.get("replacements", {}).items():
            s.placeholders[(info["entity_type"], info["canonical"])] = ph
            s.records[ph] = info
        for a in data.get("aliases", []):
            s.aliases[(a["entity_type"], a["alias"])] = a["canonical"]
        s.files_processed = list(data.get("session", {}).get("files_processed", []))
        s.file_log.update(data.get("files", {}))
        return s


# --------------------------------------------------------------------------------------
# Core engine
# --------------------------------------------------------------------------------------
class Anonymiser:
    def __init__(self, settings: dict, settings_dir: Path, session: Session, use_coref: bool = True):
        self.s = settings
        self.session = session
        self.language = settings.get("language", "en")
        self.threshold = float(settings.get("score_threshold", 0.4))
        self.allow_list = settings.get("allow_list", [])
        self.include_spacy_org = settings.get("include_spacy_organizations", False)
        self.link_partial_names = settings.get("link_partial_names", True)

        # --- Presidio analyzer with a spaCy NLP engine -------------------------------
        nlp_config = {
            "nlp_engine_name": "spacy",
            "models": [{"lang_code": self.language, "model_name": settings.get("spacy_model", "en_core_web_lg")}],
            "ner_model_configuration": {
                "model_to_presidio_entity_mapping": {
                    "PER": "PERSON", "PERSON": "PERSON", "NORP": "NRP", "FAC": "LOCATION",
                    "LOC": "LOCATION", "GPE": "LOCATION", "LOCATION": "LOCATION",
                    "ORG": "ORGANIZATION", "ORGANIZATION": "ORGANIZATION",
                    "DATE": "DATE_TIME", "TIME": "DATE_TIME",
                },
                "low_confidence_score_multiplier": 0.4,
                "low_score_entity_names": ["ORGANIZATION"],
                "labels_to_ignore": ["O", "CARDINAL", "ORDINAL", "QUANTITY", "MONEY", "PERCENT",
                                     "EVENT", "LAW", "LANGUAGE", "WORK_OF_ART", "PRODUCT"],
            },
        }
        nlp_engine = NlpEngineProvider(nlp_configuration=nlp_config).create_engine()
        self.analyzer = AnalyzerEngine(nlp_engine=nlp_engine, supported_languages=[self.language])

        for cfg in settings.get("regex_recognizers", []):
            if cfg.get("enabled", True):
                self.analyzer.registry.add_recognizer(ConfigurableRegexRecognizer(cfg, self.language))
                log.info("Added regex recogniser %s", cfg["entity"])

        self.keyword_recognizers: list[KeywordRecognizer] = []
        for cfg in settings.get("keyword_recognizers", []):
            if not cfg.get("enabled", True):
                continue
            entries = load_keyword_entries(cfg, settings_dir)
            rec = KeywordRecognizer(cfg["entity"], entries, self.language,
                                    score=float(cfg.get("score", 1.0)),
                                    case_sensitive=cfg.get("case_sensitive", False))
            self.analyzer.registry.add_recognizer(rec)
            self.keyword_recognizers.append(rec)
            log.info("Added keyword recogniser %s (%d names/aliases)", cfg["entity"], len(rec.alias_to_canonical))

        self.entities = settings.get("entities") or None
        if self.entities:  # make sure custom entities are always requested
            extra = [c["entity"] for c in settings.get("regex_recognizers", []) + settings.get("keyword_recognizers", [])
                     if c.get("enabled", True)]
            self.entities = list(dict.fromkeys([*self.entities, *extra]))

        self.anonymizer = AnonymizerEngine()

        # --- coreference ----------------------------------------------------------
        coref_cfg = settings.get("coref", {})
        self.coref_entities = set(coref_cfg.get("entities", ["PERSON", "ORGANIZATION", "ORGANIZATION_RELATED_ENTITY"]))
        self.coref_include_pronouns = coref_cfg.get("replace_pronouns", False)
        self.coref_min_chars = int(coref_cfg.get("min_chars", 40))
        self.coref: Optional[CorefResolver] = None
        if use_coref and coref_cfg.get("enabled", True):
            try:
                self.coref = CorefResolver(coref_cfg)
                log.info("Coreference model loaded (%s)", coref_cfg.get("model", "FCoref"))
            except Exception as exc:  # keep going without coref rather than fail the run
                log.warning("Coreference disabled - could not load fastcoref: %s", exc)

    # ---------------------------------------------------------------------------------
    def _source(self, r: RecognizerResult) -> str:
        name = (r.recognition_metadata or {}).get(RecognizerResult.RECOGNIZER_NAME_KEY, "")
        if name.startswith("Keyword:"):
            return "keyword"
        if name.startswith("Regex:"):
            return "regex"
        return "presidio"

    def analyse(self, text: str, use_coref: bool = True) -> list[Span]:
        if not text or not text.strip():
            return []
        results = self.analyzer.analyze(
            text=text, language=self.language, entities=self.entities,
            score_threshold=self.threshold, allow_list=self.allow_list or None,
        )
        spans = []
        for r in results:
            meta = r.recognition_metadata or {}
            # spaCy's ORG tagging is noisy - rely on the keyword list unless told otherwise
            if (r.entity_type == "ORGANIZATION" and self._source(r) == "presidio"
                    and not self.include_spacy_org):
                continue
            span = Span(r.start, r.end, r.entity_type, r.score, self._source(r), text[r.start:r.end])
            if "canonical" in meta:
                self.session.add_alias(span.entity_type, span.text, meta["canonical"])
            spans.append(span)
        spans = resolve_overlaps(spans)

        if use_coref and self.coref and len(text) >= self.coref_min_chars:
            spans = self._apply_coref(text, spans)
        if self.link_partial_names:
            self._link_partial_names(spans)
        return spans

    def _link_partial_names(self, spans: list[Span]) -> None:
        """'Okafor' or 'Mr Smith' -> same value as 'Sarah Okafor' / 'Jonathan Smith' when it
        matches exactly ONE full name already seen in this session (or this text)."""
        def tokens(s: str) -> set[str]:
            return {t for t in re.findall(r"[^\W\d_][\w'\-]*", s.casefold()) if t not in PRONOUN_LIKE}

        known = {c for (e, c) in self.session.placeholders if e == "PERSON"}
        known |= {self.session.canonical("PERSON", s.text) for s in spans if s.entity_type == "PERSON"}
        full_names = {c: tokens(c) for c in known if len(tokens(c)) >= 2}
        for s in spans:
            if s.entity_type != "PERSON":
                continue
            canon = self.session.canonical("PERSON", s.text)
            if canon in full_names:
                continue
            toks = tokens(s.text)
            if not toks:
                continue
            matches = [c for c, ft in full_names.items() if toks < ft]
            if len(matches) == 1:
                self.session.add_alias("PERSON", s.text, matches[0], self.session.alias_scope)

    def _apply_coref(self, text: str, spans: list[Span]) -> list[Span]:
        try:
            clusters = self.coref.clusters(text)
        except Exception as exc:
            log.warning("Coreference failed on a text block (%s); continuing without it", exc)
            return spans

        alias_scope = self.session.alias_scope
        added: list[Span] = []
        for cluster in clusters:
            # detected entities inside this cluster
            anchors = [s for s in spans if s.entity_type in self.coref_entities
                       and any(s.start < e and s.end > b for b, e in cluster)]
            if not anchors:
                continue
            anchor = max(anchors, key=lambda s: (SOURCE_PRIORITY.get(s.source, 0), s.end - s.start, s.score))
            # other detected mentions of the same type in the cluster share the anchor's value
            for s in anchors:
                if s is not anchor and s.entity_type == anchor.entity_type:
                    self.session.add_alias(s.entity_type, s.text, anchor.text, alias_scope)
            # undetected mentions get added as new spans
            for b, e in cluster:
                if any(b < s.end and e > s.start for s in spans + added):
                    continue
                mention = text[b:e]
                if not self.coref_include_pronouns and not looks_like_name(mention):
                    continue
                self.session.add_alias(anchor.entity_type, mention, anchor.text, alias_scope)
                added.append(Span(b, e, anchor.entity_type, anchor.score, "coref", mention))
        return resolve_overlaps(spans + added)

    # ---------------------------------------------------------------------------------
    def anonymise_spans(self, text: str, spans: list[Span]) -> tuple[str, list[tuple[Span, str]]]:
        """Run Presidio Anonymizer with a session-aware custom operator.
        Returns new text and (span, replacement) pairs in document order."""
        if not spans:
            return text, []
        session = self.session
        operators = {
            et: OperatorConfig("custom", {"lambda": (lambda value, _et=et: session.replacement(_et, value))})
            for et in {s.entity_type for s in spans}
        }
        result = self.anonymizer.anonymize(
            text=text,
            analyzer_results=[AnonymizerResult(s.entity_type, s.start, s.end, s.score) for s in spans],
            operators=operators,
        )
        items = sorted(result.items, key=lambda i: i.start)  # spans are non-overlapping -> 1:1 in order
        pairs = [(s, item.text) for s, item in zip(spans, items)]
        return result.text, pairs

    def anonymise_text(self, text: str, location: str = "", use_coref: bool = True) -> str:
        spans = self.analyse(text, use_coref=use_coref)
        new_text, pairs = self.anonymise_spans(text, spans)
        for span, repl in pairs:
            self.session.log(span, repl, location)
        return new_text


# --------------------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------------------
def read_text_file(path: Path) -> tuple[str, str]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=enc), enc
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", b"", 0, 1, f"Cannot decode {path}")


def process_csv(src: Path, dst: Path, engine: Anonymiser, cfg: dict) -> None:
    raw, enc = read_text_file(src)
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(raw.splitlines(), dialect))
    if not rows:
        dst.write_text("", encoding="utf-8")
        return

    has_header = cfg.get("has_header", True)
    header = rows[0] if has_header else [f"col_{i}" for i in range(len(rows[0]))]
    only = set(cfg.get("columns") or [])
    skip = set(cfg.get("skip_columns") or [])
    cache: dict[str, str] = {}  # identical cell values -> identical output, and much faster

    out_rows = [rows[0]] if has_header else []
    for r_idx, row in enumerate(rows[1:] if has_header else rows, start=2 if has_header else 1):
        new_row = []
        for c_idx, cell in enumerate(row):
            col = header[c_idx] if c_idx < len(header) else f"col_{c_idx}"
            if (only and col not in only) or col in skip or not cell.strip():
                new_row.append(cell)
                continue
            if cell not in cache:
                cache[cell] = engine.anonymise_text(cell, location=f"row {r_idx}, column '{col}'")
            new_row.append(cache[cell])
        out_rows.append(new_row)

    with dst.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter=dialect.delimiter, quotechar=dialect.quotechar or '"',
                            quoting=csv.QUOTE_MINIMAL)
        writer.writerows(out_rows)


# --------------------------------------------------------------------------------------
# DOCX - analyse whole document for coref context, then write back run-by-run
# so formatting (bold, fonts, etc.) is preserved.
# --------------------------------------------------------------------------------------
def _docx_parts(doc):
    from docx.oxml.ns import qn
    seen = set()
    roots = [doc.element.body]
    for section in doc.sections:
        for hf in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                   section.even_page_header, section.even_page_footer):
            try:
                if hf.is_linked_to_previous:
                    continue
                el = hf._element
            except Exception:
                continue
            if id(el) not in seen:
                seen.add(id(el))
                roots.append(el)
    for root in roots:
        for p in root.iter(qn("w:p")):
            yield p


def _paragraph_segments(p):
    """Text segments of a paragraph in order: [element|None, text]. Only w:t is editable;
    tabs/breaks are represented so offsets line up. Nested paragraphs (text boxes) are excluded."""
    from docx.oxml.ns import qn
    W_P, W_R, W_T, W_TAB, W_BR, W_CR = (qn(t) for t in ("w:p", "w:r", "w:t", "w:tab", "w:br", "w:cr"))
    segs = []
    for r in p.iter(W_R):
        owner = next(r.iterancestors(W_P), None)
        if owner is not p:
            continue
        for child in r:
            if child.tag == W_T:
                segs.append([child, child.text or ""])
            elif child.tag == W_TAB:
                segs.append([None, "\t"])
            elif child.tag in (W_BR, W_CR):
                segs.append([None, "\n"])
    return segs


def _apply_edits(segs, edits):
    """edits: sorted, non-overlapping (start, end, replacement) in paragraph coordinates."""
    # Original offsets/lengths. Edits are applied right-to-left, so text before each
    # edit's local end position is never shifted by edits already applied.
    offsets, lengths, pos = [], [], 0
    for _, t in segs:
        offsets.append(pos)
        lengths.append(len(t))
        pos += len(t)
    for start, end, repl in reversed(edits):
        placed = False
        for i, (el, t) in enumerate(segs):
            s0, s1 = offsets[i], offsets[i] + lengths[i]
            if el is None or s1 <= start or s0 >= end:
                continue
            ls, le = max(start, s0) - s0, min(end, s1) - s0
            segs[i][1] = t[:ls] + (repl if not placed else "") + t[le:]
            placed = True
    for el, t in segs:
        if el is not None and el.text != t:
            el.text = t
            el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def process_docx(src: Path, dst: Path, engine: Anonymiser, cfg: dict) -> None:
    import docx

    doc = docx.Document(str(src))
    paragraphs = [(p, _paragraph_segments(p)) for p in _docx_parts(doc)]
    paragraphs = [(p, segs) for p, segs in paragraphs if segs]

    texts = ["".join(t for _, t in segs) for _, segs in paragraphs]
    starts, pos = [], 0
    for t in texts:
        starts.append(pos)
        pos += len(t) + 1  # "\n" joiner
    full_text = "\n".join(texts)

    spans = engine.analyse(full_text, use_coref=True)
    _, pairs = engine.anonymise_spans(full_text, spans)

    edits: dict[int, list] = defaultdict(list)
    for span, repl in pairs:
        p_idx = bisect.bisect_right(starts, span.start) - 1
        engine.session.log(span, repl, location=f"paragraph {p_idx + 1}")
        first = True
        # a span may (rarely) cross a paragraph boundary - put the value in the first part
        while p_idx < len(texts) and starts[p_idx] < span.end:
            p_start, p_end = starts[p_idx], starts[p_idx] + len(texts[p_idx])
            s, e = max(span.start, p_start) - p_start, min(span.end, p_end) - p_start
            if e > s:
                edits[p_idx].append((s, e, repl if first else ""))
                first = False
            p_idx += 1

    for p_idx, p_edits in edits.items():
        _apply_edits(paragraphs[p_idx][1], sorted(p_edits))

    if cfg.get("clear_metadata", True):
        cp = doc.core_properties
        for attr in ("author", "last_modified_by", "comments", "title", "subject", "keywords", "category"):
            try:
                setattr(cp, attr, "")
            except Exception:
                pass

    doc.save(str(dst))


# --------------------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser(description="Anonymise .csv and .docx files in a folder.")
    ap.add_argument("--input", "-i", required=True, help="Folder with source files (searched recursively)")
    ap.add_argument("--output", "-o", required=True, help="Folder for anonymised files")
    ap.add_argument("--settings", "-s", default="settings.json", help="Settings JSON")
    ap.add_argument("--mapping", "-m", default=None, help="Mapping JSON to write (default: <output>/mapping.json)")
    ap.add_argument("--resume-mapping", default=None,
                    help="Existing mapping JSON - continue that session so values stay consistent")
    ap.add_argument("--no-coref", action="store_true", help="Skip coreference resolution")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")
    if not args.verbose:
        for noisy in ("presidio-analyzer", "presidio-anonymizer", "transformers", "fastcoref"):
            logging.getLogger(noisy).setLevel(logging.ERROR)

    settings_path = Path(args.settings).resolve()
    settings = json.loads(settings_path.read_text(encoding="utf-8"))
    in_dir, out_dir = Path(args.input).resolve(), Path(args.output).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    mapping_path = Path(args.mapping) if args.mapping else out_dir / "mapping.json"

    fmt = settings.get("replacement_format", "<{entity_type}_{index}>")
    alias_scope = settings.get("coref", {}).get("alias_scope", "session")
    if args.resume_mapping:
        session = Session.from_json(json.loads(Path(args.resume_mapping).read_text(encoding="utf-8")), fmt, alias_scope)
        log.info("Resumed session from %s (%d placeholders)", args.resume_mapping, len(session.placeholders))
    else:
        session = Session(fmt, alias_scope)

    engine = Anonymiser(settings, settings_path.parent, session, use_coref=not args.no_coref)

    files = sorted(p for p in in_dir.rglob("*")
                   if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES and not p.name.startswith("~$"))
    if not files:
        log.warning("No .csv or .docx files found in %s", in_dir)

    for src in files:
        rel = src.relative_to(in_dir)
        dst = out_dir / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        session.start_file(str(rel))
        before = len(session.file_log.get(str(rel), []))
        try:
            if src.suffix.lower() == ".csv":
                process_csv(src, dst, engine, settings.get("csv", {}))
            else:
                process_docx(src, dst, engine, settings.get("docx", {}))
            log.info("%s -> %d replacements", rel, len(session.file_log.get(str(rel), [])) - before)
        except Exception:
            log.exception("Failed to process %s", rel)

    mapping_path.parent.mkdir(parents=True, exist_ok=True)
    mapping_path.write_text(json.dumps(session.to_json(str(in_dir), str(settings_path)), indent=2,
                                       ensure_ascii=False), encoding="utf-8")
    log.info("Mapping written to %s (%d unique values). It contains the original PII - store it securely.",
             mapping_path, len(session.records))


if __name__ == "__main__":
    main()
